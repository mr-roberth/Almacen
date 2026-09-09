#!/usr/bin/env python3
"""Genera la migración 008 desde los cuatro catálogos depurados de Oleolab.

El generador no se conecta a MySQL y no modifica los Excel. La migración resultante
es acumulativa, conserva transacciones anteriores y registra incompatibilidades de
origen para que nunca se conviertan silenciosamente en datos maestros.
"""

from __future__ import annotations

import argparse
import calendar
import collections
import hashlib
import math
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl


FORMULATION_CODES = {
    "CMP04038SX BG": "CMP04038SX-BG",
    "CMP04011": "CMP04011",
    "CMP04003CX CORAL": "CMP04003CX-CORAL",
    "CMP04012": "CMP04012",
    "CMP04001": "CMP04001",
    "CMP04003CX SOVENA": "CMP04003CX-SOVENA",
    "CMP04003SX LILYS": "CMP04003SX-LILYS",
    "CMP04003SX PK SX": "CMP04003SX-PK-SX",
    "CMP04003CX REFINADOS": "CMP04003CX-REFINADOS",
    "Prensado en frio CMP04001": "CMP04001-PRENSADO-FRIO",
    "Aceite de uva": "FORM-ACEITE-UVA",
    "Blend Valley Foods": "FORM-BLEND-VALLEY",
}
FORMULA_BLOCKS = ((2, 1), (2, 8), (2, 14), (13, 1), (13, 8), (13, 14),
                  (24, 1), (24, 8), (24, 14), (35, 1), (35, 8), (35, 14))
MONTHS = tuple(range(1, 13))


def text(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def code(value: object) -> str:
    return text(value).upper()


def positive_number(value: object, label: str, row: int) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Fila {row}: {label} no es numérico.") from exc
    if not math.isfinite(result) or result <= 0:
        raise ValueError(f"Fila {row}: {label} debe ser mayor que cero.")
    return result


def sql_value(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            raise ValueError("No se puede generar SQL con un número no finito.")
        return f"{float(value):.8f}".rstrip("0").rstrip(".") or "0"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def values_sql(table: str, columns: str, rows: list[tuple[object, ...]], chunk_size: int = 200) -> str:
    statements: list[str] = []
    for start in range(0, len(rows), chunk_size):
        group = rows[start:start + chunk_size]
        encoded = ["(" + ",".join(sql_value(value) for value in row) + ")" for row in group]
        statements.append(f"INSERT INTO {table} ({columns}) VALUES\n  " + ",\n  ".join(encoded) + ";")
    return "\n\n".join(statements)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def load_products(path: Path) -> tuple[collections.OrderedDict[str, dict[str, object]], list[dict[str, object]]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheet = workbook["CAT_Producto terminado"]
    expected = ["Codigo_PT", "Nombe del producto", "Presentación ml"]
    if [text(sheet.cell(1, column).value) for column in range(1, 4)] != expected:
        raise ValueError("Los encabezados del catálogo de producto terminado cambiaron.")
    products: collections.OrderedDict[str, dict[str, object]] = collections.OrderedDict()
    for row in range(2, sheet.max_row + 1):
        sku = code(sheet.cell(row, 1).value)
        if not sku:
            continue
        if not re.fullmatch(r"PT\d{5,}", sku):
            raise ValueError(f"Fila {row}: código de producto terminado no válido: {sku}.")
        if sku in products:
            raise ValueError(f"Fila {row}: código de producto terminado repetido: {sku}.")
        products[sku] = {
            "name": text(sheet.cell(row, 2).value),
            "ml": positive_number(sheet.cell(row, 3).value, "Presentación ml", row),
            "row": row,
        }
    if not products:
        raise ValueError("El catálogo de producto terminado está vacío.")

    users_sheet = workbook["CAT_Usuarios"]
    users = []
    for row in range(2, users_sheet.max_row + 1):
        email = text(users_sheet.cell(row, 6).value).lower()
        if email:
            users.append({"row": row, "email": email, "name": text(users_sheet.cell(row, 2).value)})
    return products, users


def load_packaging(path: Path, products: collections.OrderedDict[str, dict[str, object]]):
    formula_sheet = openpyxl.load_workbook(path, data_only=False, read_only=False).active
    value_sheet = openpyxl.load_workbook(path, data_only=True, read_only=False).active
    expected = [
        "Codigo_PT", "Nombre terminado a fabricar", "Codigo de material de empaque",
        "Nombre de material de empaque", "Cajas x Pallet", "Piezas x Caja",
        "Cantidad x caja", "Catidad por pieza", "Unidad",
    ]
    if [text(formula_sheet.cell(1, column).value) for column in range(1, 10)] != expected:
        raise ValueError("Los encabezados del BOM de materiales de empaque cambiaron.")

    grouped: collections.OrderedDict[str, list[dict[str, object]]] = collections.OrderedDict()
    material_names: dict[str, collections.Counter[str]] = collections.defaultdict(collections.Counter)
    material_units: dict[str, collections.Counter[str]] = collections.defaultdict(collections.Counter)
    current = ""
    issues: list[tuple[str, str, int, str, str]] = []
    seen_headers: set[str] = set()
    for row in range(2, formula_sheet.max_row + 1):
        product_code = code(formula_sheet.cell(row, 1).value)
        material_code = code(formula_sheet.cell(row, 3).value)
        if not product_code and not material_code:
            continue
        if product_code == material_code and product_code.startswith("PT"):
            current = product_code
            if current in seen_headers:
                raise ValueError(f"Fila {row}: la lista {current} aparece dos veces.")
            seen_headers.add(current)
            grouped[current] = []
            continue
        if not current:
            raise ValueError(f"Fila {row}: componente sin producto terminado anterior.")
        if product_code and product_code != current:
            issues.append((
                "WARNING",
                "PACKAGING_BLOCK_CODE_MISMATCH",
                row,
                material_code,
                f"El renglón indicaba {product_code}, pero pertenece al bloque {current}; se conservó {current}.",
            ))
        if not material_code.startswith("CME"):
            raise ValueError(f"Fila {row}: el material {material_code} no tiene prefijo CME.")
        name = text(formula_sheet.cell(row, 4).value)
        raw_unit = code(formula_sheet.cell(row, 9).value)
        unit = "KG" if raw_unit == "KG" else "EA"
        pieces = positive_number(formula_sheet.cell(row, 6).value, "Piezas x caja", row)
        per_box = positive_number(formula_sheet.cell(row, 7).value, "Cantidad x caja", row)
        cached = value_sheet.cell(row, 8).value
        quantity = float(cached) if isinstance(cached, (int, float)) and cached > 0 else per_box / pieces
        grouped[current].append({"code": material_code, "name": name, "unit": unit, "qty": quantity, "row": row})
        material_names[material_code][name] += 1
        material_units[material_code][unit] += 1

    missing_bom = sorted(set(products) - set(grouped))
    if missing_bom:
        raise ValueError("Faltan listas de empaque para: " + ", ".join(missing_bom[:10]))

    official_grouped = collections.OrderedDict((sku, grouped[sku]) for sku in products)
    materials = []
    for material_code in sorted(material_names):
        names = material_names[material_code]
        units = material_units[material_code]
        canonical_name = sorted(names.items(), key=lambda item: (-item[1], item[0]))[0][0]
        canonical_unit = sorted(units.items(), key=lambda item: (-item[1], item[0]))[0][0]
        first_row = min(component["row"] for rows in grouped.values() for component in rows if component["code"] == material_code)
        materials.append((material_code, canonical_name, canonical_unit, first_row))
        if len(names) > 1:
            issues.append(("WARNING", "PACKAGING_NAME_CONFLICT", first_row, material_code,
                           "La clave tiene variantes de nombre; se usó la variante más frecuente: " + canonical_name))
        if len(units) > 1:
            issues.append(("WARNING", "PACKAGING_UOM_CONFLICT", first_row, material_code,
                           "La clave aparece con unidades distintas; se usó la unidad más frecuente: " + canonical_unit))
    return official_grouped, materials, issues


def load_formulas(path: Path):
    sheet = openpyxl.load_workbook(path, data_only=True, read_only=False).active
    formulas = []
    raw_materials: dict[str, str] = {}
    issues: list[tuple[str, str, int, str, str]] = []
    for header_row, start_col in FORMULA_BLOCKS:
        raw_header = text(sheet.cell(header_row, start_col).value)
        label = text(raw_header.split("=", 1)[1] if "=" in raw_header else raw_header)
        if label not in FORMULATION_CODES:
            raise ValueError(f"No existe una clave normalizada para la formulación: {label}.")
        formula_code = FORMULATION_CODES[label]
        components = []
        for row in range(header_row + 2, min(header_row + 9, sheet.max_row + 1)):
            component_code = code(sheet.cell(row, start_col).value)
            component_name = text(sheet.cell(row, start_col + 1).value)
            percent_value = sheet.cell(row, start_col + 2).value
            percent = float(percent_value) if isinstance(percent_value, (int, float)) else 0.0
            if not component_code and component_name.lower() == "aceite italia":
                component_code = "CMP03008"
                issues.append(("WARNING", "GENERATED_RAW_CODE", row, component_code,
                               "El archivo no incluía clave para Aceite Italia; se asignó CMP03008."))
            if component_code and component_name:
                raw_materials[component_code] = component_name
            if percent > 0:
                if not component_code or not component_name:
                    raise ValueError(f"Fila {row}: componente de formulación incompleto.")
                components.append((component_code, component_name, percent, row))
        total = sum(component[2] for component in components)
        if not math.isclose(total, 100.0, abs_tol=0.0001):
            raise ValueError(f"La formulación {formula_code} suma {total}%, no 100%.")
        if not re.search(r"CMP\d", label, flags=re.IGNORECASE):
            issues.append(("WARNING", "GENERATED_FORMULA_CODE", header_row, formula_code,
                           f"La formulación '{label}' no tenía código; se asignó {formula_code}."))
        formulas.append({"code": formula_code, "name": label, "components": components, "row": header_row})
    return formulas, raw_materials, issues


def allocate_month(year: int, month: int, quantity: float) -> list[tuple[date, float]]:
    """Distribuye botellas por semana según días laborables, conservando el total."""
    workdays: dict[date, int] = collections.OrderedDict()
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        current = date(year, month, day)
        if current.weekday() < 5:
            monday = current - timedelta(days=current.weekday())
            workdays[monday] = workdays.get(monday, 0) + 1
    if not workdays or quantity <= 0:
        return []
    if math.isclose(quantity, round(quantity), abs_tol=0.000001):
        total = int(round(quantity))
        exact = [(week, total * days / sum(workdays.values())) for week, days in workdays.items()]
        allocated = [(week, math.floor(value)) for week, value in exact]
        remainder = total - sum(value for _, value in allocated)
        order = sorted(range(len(exact)), key=lambda index: (-(exact[index][1] - math.floor(exact[index][1])), exact[index][0]))
        values = [value for _, value in allocated]
        for index in order[:remainder]:
            values[index] += 1
        return [(exact[index][0], float(values[index])) for index in range(len(exact)) if values[index] > 0]
    total_days = sum(workdays.values())
    return [(week, quantity * days / total_days) for week, days in workdays.items()]


def load_forecast(path: Path, products: collections.OrderedDict[str, dict[str, object]], year: int):
    sheet = openpyxl.load_workbook(path, data_only=True, read_only=False)["FORECAST_PLAN"]
    expected = ["Código Interno", "Descripción", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    if [text(sheet.cell(1, column).value) for column in range(1, 15)] != expected:
        raise ValueError("Los encabezados del Forecast cambiaron.")
    monthly = []
    weekly = []
    issues: list[tuple[str, str, int, str, str]] = []
    for row in range(2, sheet.max_row + 1):
        sku = code(sheet.cell(row, 1).value)
        if not sku:
            continue
        if sku not in products:
            issues.append(("ERROR", "FORECAST_UNKNOWN_PRODUCT", row, sku,
                           "La clave del Forecast no existe en el catálogo oficial de producto terminado y no se importó."))
            continue
        for month in MONTHS:
            value = sheet.cell(row, month + 2).value
            if value in (None, ""):
                issues.append(("WARNING", "FORECAST_MISSING_VALUE", row, sku,
                               f"El mes {month} no tiene cantidad; se dejó sin dato, no se convirtió en cero."))
                continue
            try:
                quantity = float(value)
            except (TypeError, ValueError):
                issues.append(("ERROR", "FORECAST_INVALID_VALUE", row, sku,
                               f"El mes {month} no contiene una cantidad numérica y no se importó."))
                continue
            if not math.isfinite(quantity) or quantity < 0:
                issues.append(("ERROR", "FORECAST_INVALID_VALUE", row, sku,
                               f"El mes {month} contiene una cantidad negativa o no finita y no se importó."))
                continue
            period = date(year, month, 1)
            monthly.append((sku, period.isoformat(), quantity, row))
            weekly.extend((sku, week.isoformat(), amount, row, month) for week, amount in allocate_month(year, month, quantity))
    return monthly, weekly, issues


def build_sql(paths: list[Path], products, users, packaging, materials, formulas, raw_materials,
              monthly, weekly, issues, year: int) -> str:
    combined_hash = hashlib.sha256("".join(sha256(path) for path in paths).encode()).hexdigest().upper()
    packaging_hash = sha256(paths[3])
    formula_hash = sha256(paths[2])
    forecast_hash = sha256(paths[1])
    source_ref = f"CORE-{combined_hash[:16]}"

    product_rows = [(sku, data["name"], data["ml"], data["row"]) for sku, data in products.items()]
    component_rows = []
    for sku, rows in packaging.items():
        for line_no, component in enumerate(rows, start=1):
            component_rows.append((sku, line_no, component["code"], component["qty"], component["row"]))
    raw_rows = [("CMP01001", "Aguacate (fruta)", "RAW_FRUIT", "KG", 0)]
    raw_rows.extend((sku, name, "BULK_OIL", "L", 1) for sku, name in sorted(raw_materials.items()))
    formula_rows = [(formula["code"], formula["name"], formula["row"]) for formula in formulas]
    formula_component_rows = []
    for formula in formulas:
        for line_no, component in enumerate(formula["components"], start=1):
            formula_component_rows.append((formula["code"], line_no, component[0], component[2], component[3]))
    issue_rows = [(source_ref, severity, issue_code, entity_code, row_number, message)
                  for severity, issue_code, row_number, entity_code, message in issues]
    issue_rows.append((source_ref, "WARNING", "FORMULA_ASSIGNMENT_REQUIRED", None, None,
                       "El archivo de producto terminado no indica qué formulación corresponde a cada clave. Debe confirmarse en el módulo Listas de materiales antes de calcular aceites."))
    issue_rows.append((source_ref, "INFO", "USERS_REVIEWED_NOT_PUBLISHED", None, None,
                       f"Se validaron {len(users)} usuarios del archivo privado. Sus correos y teléfonos no se incluyen en la migración ni en GitHub."))

    sections = [f"""-- Oleolab Almacenes · Migración 008
-- Reingeniería del núcleo: catálogos oficiales, listas de materiales, Forecast semanal y control de saldos iniciales.
-- Generada desde cuatro archivos Excel. Fuente combinada: {combined_hash}
-- Importar con la base de datos seleccionada y después de migration_v1_7.sql.
-- Es acumulativa: no borra usuarios, existencias, movimientos ni órdenes anteriores.

SET NAMES utf8mb4;
SET time_zone = '+00:00';

CREATE TABLE IF NOT EXISTS product_presentations (
  finished_item_id BIGINT UNSIGNED NOT NULL,
  presentation_ml DECIMAL(12,3) NOT NULL,
  source_reference VARCHAR(100) NULL,
  updated_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6) ON UPDATE CURRENT_TIMESTAMP(6),
  PRIMARY KEY (finished_item_id),
  CONSTRAINT fk_product_presentations_item FOREIGN KEY (finished_item_id) REFERENCES items(id) ON DELETE CASCADE,
  CONSTRAINT chk_product_presentations_ml CHECK (presentation_ml > 0)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS oil_formulations (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  formulation_code VARCHAR(80) NOT NULL,
  name VARCHAR(190) NOT NULL,
  is_active TINYINT(1) NOT NULL DEFAULT 1,
  created_by BIGINT UNSIGNED NULL,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  PRIMARY KEY (id),
  UNIQUE KEY uq_oil_formulations_code (formulation_code),
  CONSTRAINT fk_oil_formulations_user FOREIGN KEY (created_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_oil_formulations_active CHECK (is_active IN (0,1))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS oil_formulation_versions (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  oil_formulation_id BIGINT UNSIGNED NOT NULL,
  version_no INT UNSIGNED NOT NULL,
  effective_from DATE NOT NULL,
  effective_to DATE NULL,
  status VARCHAR(20) NOT NULL DEFAULT 'DRAFT',
  source_reference VARCHAR(100) NULL,
  approved_by BIGINT UNSIGNED NULL,
  approved_at DATETIME(6) NULL,
  PRIMARY KEY (id),
  UNIQUE KEY uq_oil_formula_version (oil_formulation_id,version_no),
  UNIQUE KEY uq_oil_formula_source (oil_formulation_id,source_reference),
  CONSTRAINT fk_oil_formula_version_formula FOREIGN KEY (oil_formulation_id) REFERENCES oil_formulations(id) ON DELETE CASCADE,
  CONSTRAINT fk_oil_formula_version_user FOREIGN KEY (approved_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_oil_formula_version_dates CHECK (effective_to IS NULL OR effective_to >= effective_from),
  CONSTRAINT chk_oil_formula_version_status CHECK (status IN ('DRAFT','ACTIVE','OBSOLETE'))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS oil_formulation_components (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  oil_formulation_version_id BIGINT UNSIGNED NOT NULL,
  line_no INT UNSIGNED NOT NULL,
  raw_item_id BIGINT UNSIGNED NOT NULL,
  percentage DECIMAL(9,6) NOT NULL,
  PRIMARY KEY (id),
  UNIQUE KEY uq_oil_formula_component_line (oil_formulation_version_id,line_no),
  UNIQUE KEY uq_oil_formula_component_item (oil_formulation_version_id,raw_item_id),
  CONSTRAINT fk_oil_formula_component_version FOREIGN KEY (oil_formulation_version_id) REFERENCES oil_formulation_versions(id) ON DELETE CASCADE,
  CONSTRAINT fk_oil_formula_component_item FOREIGN KEY (raw_item_id) REFERENCES items(id) ON DELETE RESTRICT,
  CONSTRAINT chk_oil_formula_component_pct CHECK (percentage > 0 AND percentage <= 100)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS finished_product_bom_assignments (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  finished_item_id BIGINT UNSIGNED NOT NULL,
  packaging_bom_version_id BIGINT UNSIGNED NOT NULL,
  oil_formula_version_id BIGINT UNSIGNED NULL,
  compiled_bom_version_id BIGINT UNSIGNED NULL,
  effective_from DATE NOT NULL,
  effective_to DATE NULL,
  assignment_status VARCHAR(30) NOT NULL DEFAULT 'PENDING_FORMULATION',
  change_reason VARCHAR(500) NULL,
  created_by BIGINT UNSIGNED NULL,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  PRIMARY KEY (id),
  UNIQUE KEY uq_finished_product_assignment_date (finished_item_id,effective_from),
  KEY ix_finished_product_assignment_active (finished_item_id,effective_from,effective_to),
  CONSTRAINT fk_finished_assignment_item FOREIGN KEY (finished_item_id) REFERENCES items(id) ON DELETE RESTRICT,
  CONSTRAINT fk_finished_assignment_packaging FOREIGN KEY (packaging_bom_version_id) REFERENCES bom_versions(id) ON DELETE RESTRICT,
  CONSTRAINT fk_finished_assignment_formula FOREIGN KEY (oil_formula_version_id) REFERENCES oil_formulation_versions(id) ON DELETE RESTRICT,
  CONSTRAINT fk_finished_assignment_compiled FOREIGN KEY (compiled_bom_version_id) REFERENCES bom_versions(id) ON DELETE RESTRICT,
  CONSTRAINT fk_finished_assignment_user FOREIGN KEY (created_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_finished_assignment_dates CHECK (effective_to IS NULL OR effective_to >= effective_from),
  CONSTRAINT chk_finished_assignment_status CHECK (assignment_status IN ('PENDING_FORMULATION','ACTIVE','OBSOLETE'))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS forecast_weekly_lines (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  forecast_version_id BIGINT UNSIGNED NOT NULL,
  finished_item_id BIGINT UNSIGNED NOT NULL,
  week_start DATE NOT NULL,
  forecast_qty DECIMAL(18,6) NOT NULL,
  allocation_method VARCHAR(40) NOT NULL DEFAULT 'WORKDAYS',
  source_month TINYINT UNSIGNED NOT NULL,
  source_row INT UNSIGNED NULL,
  PRIMARY KEY (id),
  UNIQUE KEY uq_forecast_weekly_dimension (forecast_version_id,finished_item_id,week_start,source_month),
  KEY ix_forecast_weekly_period (week_start,finished_item_id),
  CONSTRAINT fk_forecast_weekly_version FOREIGN KEY (forecast_version_id) REFERENCES forecast_versions(id) ON DELETE CASCADE,
  CONSTRAINT fk_forecast_weekly_item FOREIGN KEY (finished_item_id) REFERENCES items(id) ON DELETE RESTRICT,
  CONSTRAINT chk_forecast_weekly_qty CHECK (forecast_qty >= 0),
  CONSTRAINT chk_forecast_weekly_month CHECK (source_month BETWEEN 1 AND 12)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS warehouse_inventory_settings (
  warehouse_id BIGINT UNSIGNED NOT NULL,
  initial_capture_enabled TINYINT(1) NOT NULL DEFAULT 1,
  disabled_at DATETIME(6) NULL,
  disabled_by BIGINT UNSIGNED NULL,
  PRIMARY KEY (warehouse_id),
  CONSTRAINT fk_warehouse_inventory_setting_warehouse FOREIGN KEY (warehouse_id) REFERENCES warehouses(id) ON DELETE CASCADE,
  CONSTRAINT fk_warehouse_inventory_setting_user FOREIGN KEY (disabled_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_warehouse_inventory_setting_enabled CHECK (initial_capture_enabled IN (0,1))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS source_import_issues (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  source_reference VARCHAR(100) NOT NULL,
  severity VARCHAR(10) NOT NULL,
  issue_code VARCHAR(60) NOT NULL,
  entity_code VARCHAR(100) NULL,
  source_row INT UNSIGNED NULL,
  message VARCHAR(1000) NOT NULL,
  is_resolved TINYINT(1) NOT NULL DEFAULT 0,
  resolved_by BIGINT UNSIGNED NULL,
  resolved_at DATETIME(6) NULL,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  PRIMARY KEY (id),
  KEY ix_source_issues_open (source_reference,is_resolved,severity),
  CONSTRAINT fk_source_issues_user FOREIGN KEY (resolved_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_source_issues_severity CHECK (severity IN ('INFO','WARNING','ERROR')),
  CONSTRAINT chk_source_issues_resolved CHECK (is_resolved IN (0,1))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

CREATE TABLE IF NOT EXISTS planned_material_arrivals (
  id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
  arrival_number VARCHAR(60) NOT NULL,
  item_id BIGINT UNSIGNED NOT NULL,
  warehouse_id BIGINT UNSIGNED NOT NULL,
  planned_week DATE NOT NULL,
  expected_qty DECIMAL(18,6) NOT NULL,
  uom_id BIGINT UNSIGNED NOT NULL,
  supplier_id BIGINT UNSIGNED NULL,
  status VARCHAR(25) NOT NULL DEFAULT 'PLANNED',
  notes VARCHAR(1000) NULL,
  inbound_appointment_id BIGINT UNSIGNED NULL,
  created_by BIGINT UNSIGNED NULL,
  created_at DATETIME(6) NOT NULL DEFAULT CURRENT_TIMESTAMP(6),
  PRIMARY KEY (id),
  UNIQUE KEY uq_planned_arrival_number (arrival_number),
  KEY ix_planned_arrival_week_item (planned_week,item_id,status),
  CONSTRAINT fk_planned_arrival_item FOREIGN KEY (item_id) REFERENCES items(id) ON DELETE RESTRICT,
  CONSTRAINT fk_planned_arrival_warehouse FOREIGN KEY (warehouse_id) REFERENCES warehouses(id) ON DELETE RESTRICT,
  CONSTRAINT fk_planned_arrival_uom FOREIGN KEY (uom_id) REFERENCES units_of_measure(id) ON DELETE RESTRICT,
  CONSTRAINT fk_planned_arrival_supplier FOREIGN KEY (supplier_id) REFERENCES suppliers(id) ON DELETE SET NULL,
  CONSTRAINT fk_planned_arrival_appointment FOREIGN KEY (inbound_appointment_id) REFERENCES inbound_appointments(id) ON DELETE SET NULL,
  CONSTRAINT fk_planned_arrival_user FOREIGN KEY (created_by) REFERENCES app_users(id) ON DELETE SET NULL,
  CONSTRAINT chk_planned_arrival_qty CHECK (expected_qty > 0),
  CONSTRAINT chk_planned_arrival_status CHECK (status IN ('PLANNED','NOTIFIED','APPOINTMENT_CREATED','RECEIVED','CANCELLED'))
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;

INSERT IGNORE INTO app_modules (code,name,route_path,sort_order) VALUES
('BOM_MANAGEMENT','Listas de materiales','/planeacion/listas-materiales',45),
('PLANNED_ARRIVALS','Llegadas planeadas','/planeacion/llegadas',55);

INSERT IGNORE INTO permissions (module_id,action_code,description)
SELECT m.id,a.action_code,CONCAT(a.action_code,' en ',m.name)
FROM app_modules m CROSS JOIN (
  SELECT 'VIEW' action_code UNION ALL SELECT 'CREATE' UNION ALL SELECT 'UPDATE' UNION ALL
  SELECT 'APPROVE' UNION ALL SELECT 'REJECT' UNION ALL SELECT 'POST' UNION ALL
  SELECT 'EXPORT' UNION ALL SELECT 'ADMIN'
) a WHERE m.code IN ('BOM_MANAGEMENT','PLANNED_ARRIVALS');

INSERT IGNORE INTO role_permissions (role_id,permission_id)
SELECT r.id,p.id FROM roles r CROSS JOIN permissions p WHERE r.code='SUPER_ADMIN';

SET @column_exists=(SELECT COUNT(*) FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name='production_orders' AND column_name='bom_assignment_id');
SET @ddl=IF(@column_exists=0,'ALTER TABLE production_orders ADD COLUMN bom_assignment_id BIGINT UNSIGNED NULL AFTER bom_version_id','SELECT 1');
PREPARE migration_statement FROM @ddl; EXECUTE migration_statement; DEALLOCATE PREPARE migration_statement;
SET @column_exists=(SELECT COUNT(*) FROM information_schema.columns WHERE table_schema=DATABASE() AND table_name='production_orders' AND column_name='oil_formula_version_id');
SET @ddl=IF(@column_exists=0,'ALTER TABLE production_orders ADD COLUMN oil_formula_version_id BIGINT UNSIGNED NULL AFTER bom_assignment_id','SELECT 1');
PREPARE migration_statement FROM @ddl; EXECUTE migration_statement; DEALLOCATE PREPARE migration_statement;

SET @fk_exists=(SELECT COUNT(*) FROM information_schema.referential_constraints WHERE constraint_schema=DATABASE() AND constraint_name='fk_production_order_assignment');
SET @ddl=IF(@fk_exists=0,'ALTER TABLE production_orders ADD CONSTRAINT fk_production_order_assignment FOREIGN KEY (bom_assignment_id) REFERENCES finished_product_bom_assignments(id) ON DELETE SET NULL','SELECT 1');
PREPARE migration_statement FROM @ddl; EXECUTE migration_statement; DEALLOCATE PREPARE migration_statement;
SET @fk_exists=(SELECT COUNT(*) FROM information_schema.referential_constraints WHERE constraint_schema=DATABASE() AND constraint_name='fk_production_order_oil_formula');
SET @ddl=IF(@fk_exists=0,'ALTER TABLE production_orders ADD CONSTRAINT fk_production_order_oil_formula FOREIGN KEY (oil_formula_version_id) REFERENCES oil_formulation_versions(id) ON DELETE SET NULL','SELECT 1');
PREPARE migration_statement FROM @ddl; EXECUTE migration_statement; DEALLOCATE PREPARE migration_statement;

START TRANSACTION;

SET @system_admin_id=(SELECT ur.user_id FROM user_roles ur INNER JOIN roles r ON r.id=ur.role_id INNER JOIN app_users u ON u.id=ur.user_id WHERE ur.is_active=1 AND r.code='SUPER_ADMIN' AND u.status='ACTIVE' ORDER BY ur.id LIMIT 1);
SET @source_reference='{source_ref}';
SET @packaging_reference='CORE-ME-{packaging_hash[:16]}';
SET @formula_reference='CORE-OIL-{formula_hash[:16]}';
SET @forecast_hash='{forecast_hash}';

INSERT INTO warehouses (code,name,warehouse_type,area_id,is_active)
SELECT source.code,source.name,source.warehouse_type,a.id,1 FROM (
  SELECT 'MP' code,'Almacén de materias primas' name,'RAW_MATERIAL' warehouse_type UNION ALL
  SELECT 'EMPAQUE','Almacén de materiales de empaque','PACKAGING' UNION ALL
  SELECT 'PT','Almacén de producto terminado','FINISHED_GOODS'
) source INNER JOIN areas a ON a.code='ALMACEN'
ON DUPLICATE KEY UPDATE name=VALUES(name),warehouse_type=VALUES(warehouse_type),area_id=VALUES(area_id),is_active=1;

UPDATE warehouse_locations l INNER JOIN warehouses w ON w.id=l.warehouse_id
SET l.name=CASE l.location_type WHEN 'RECEIVING' THEN 'Recepción pendiente de ingreso a Calidad' WHEN 'QUALITY_HOLD' THEN 'Pendiente de liberación por Calidad' WHEN 'STORAGE' THEN 'Existencia disponible' WHEN 'REJECTED' THEN 'Producto rechazado · no disponible' ELSE l.name END
WHERE w.code IN ('MP','EMPAQUE','PT');

INSERT INTO warehouse_locations (warehouse_id,code,name,location_type,is_blocked,is_active)
SELECT w.id,'WID','Pendiente de liberación por Calidad','QUALITY_HOLD',1,1 FROM warehouses w
WHERE w.code IN ('MP','EMPAQUE','PT') AND NOT EXISTS (SELECT 1 FROM warehouse_locations l WHERE l.warehouse_id=w.id AND l.location_type='QUALITY_HOLD');
INSERT INTO warehouse_locations (warehouse_id,code,name,location_type,is_blocked,is_active)
SELECT w.id,'DISPONIBLE','Existencia disponible','STORAGE',0,1 FROM warehouses w
WHERE w.code IN ('MP','EMPAQUE','PT') AND NOT EXISTS (SELECT 1 FROM warehouse_locations l WHERE l.warehouse_id=w.id AND l.location_type='STORAGE');
INSERT INTO warehouse_locations (warehouse_id,code,name,location_type,is_blocked,is_active)
SELECT w.id,'RECHAZADO','Producto rechazado · no disponible','REJECTED',1,1 FROM warehouses w
WHERE w.code IN ('MP','EMPAQUE','PT') AND NOT EXISTS (SELECT 1 FROM warehouse_locations l WHERE l.warehouse_id=w.id AND l.location_type='REJECTED');

INSERT IGNORE INTO warehouse_inventory_settings (warehouse_id,initial_capture_enabled)
SELECT id,1 FROM warehouses WHERE code IN ('MP','EMPAQUE','PT');

DROP TEMPORARY TABLE IF EXISTS tmp_core_pt;
DROP TEMPORARY TABLE IF EXISTS tmp_core_materials;
DROP TEMPORARY TABLE IF EXISTS tmp_core_pack_components;
DROP TEMPORARY TABLE IF EXISTS tmp_core_raw;
DROP TEMPORARY TABLE IF EXISTS tmp_core_formulations;
DROP TEMPORARY TABLE IF EXISTS tmp_core_formula_components;
DROP TEMPORARY TABLE IF EXISTS tmp_core_forecast_monthly;
DROP TEMPORARY TABLE IF EXISTS tmp_core_forecast_weekly;
DROP TEMPORARY TABLE IF EXISTS tmp_core_issues;
CREATE TEMPORARY TABLE tmp_core_pt (product_code VARCHAR(80) PRIMARY KEY,product_name VARCHAR(190),presentation_ml DECIMAL(12,3),source_row INT UNSIGNED) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_materials (material_code VARCHAR(80) PRIMARY KEY,material_name VARCHAR(190),uom_code VARCHAR(20),source_row INT UNSIGNED) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_pack_components (product_code VARCHAR(80),line_no INT UNSIGNED,material_code VARCHAR(80),qty_per_unit DECIMAL(18,8),source_row INT UNSIGNED,PRIMARY KEY(product_code,line_no),UNIQUE KEY uq_tmp_pack_component(product_code,material_code)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_raw (material_code VARCHAR(80) PRIMARY KEY,material_name VARCHAR(190),item_type VARCHAR(30),uom_code VARCHAR(20),source_row INT UNSIGNED) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_formulations (formulation_code VARCHAR(80) PRIMARY KEY,formulation_name VARCHAR(190),source_row INT UNSIGNED) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_formula_components (formulation_code VARCHAR(80),line_no INT UNSIGNED,material_code VARCHAR(80),percentage DECIMAL(9,6),source_row INT UNSIGNED,PRIMARY KEY(formulation_code,line_no),UNIQUE KEY uq_tmp_formula_component(formulation_code,material_code)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_forecast_monthly (product_code VARCHAR(80),period_start DATE,forecast_qty DECIMAL(18,6),source_row INT UNSIGNED,PRIMARY KEY(product_code,period_start)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_forecast_weekly (product_code VARCHAR(80),week_start DATE,forecast_qty DECIMAL(18,6),source_row INT UNSIGNED,source_month TINYINT UNSIGNED,PRIMARY KEY(product_code,week_start,source_month)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;
CREATE TEMPORARY TABLE tmp_core_issues (source_reference VARCHAR(100),severity VARCHAR(10),issue_code VARCHAR(60),entity_code VARCHAR(100),source_row INT UNSIGNED,message VARCHAR(1000)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_0900_ai_ci;""",
        values_sql("tmp_core_pt", "product_code,product_name,presentation_ml,source_row", product_rows),
        values_sql("tmp_core_materials", "material_code,material_name,uom_code,source_row", materials),
        values_sql("tmp_core_pack_components", "product_code,line_no,material_code,qty_per_unit,source_row", component_rows),
        values_sql("tmp_core_raw", "material_code,material_name,item_type,uom_code,source_row", raw_rows),
        values_sql("tmp_core_formulations", "formulation_code,formulation_name,source_row", formula_rows),
        values_sql("tmp_core_formula_components", "formulation_code,line_no,material_code,percentage,source_row", formula_component_rows),
        values_sql("tmp_core_forecast_monthly", "product_code,period_start,forecast_qty,source_row", monthly),
        values_sql("tmp_core_forecast_weekly", "product_code,week_start,forecast_qty,source_row,source_month", weekly),
        values_sql("tmp_core_issues", "source_reference,severity,issue_code,entity_code,source_row,message", issue_rows),
        f"""
-- Catálogos oficiales. Los identificadores existentes se conservan.
INSERT INTO items (category_id,base_uom_id,sku,name,item_type,lot_controlled,quality_required,created_by)
SELECT c.id,u.id,p.product_code,p.product_name,'FINISHED_GOOD',1,1,@system_admin_id
FROM tmp_core_pt p INNER JOIN item_categories c ON c.code='FG' INNER JOIN units_of_measure u ON u.code='EA'
ON DUPLICATE KEY UPDATE category_id=VALUES(category_id),base_uom_id=items.base_uom_id,name=VALUES(name),item_type='FINISHED_GOOD',lot_controlled=1,quality_required=1,is_active=1;

INSERT INTO product_presentations (finished_item_id,presentation_ml,source_reference)
SELECT i.id,p.presentation_ml,@source_reference FROM tmp_core_pt p INNER JOIN items i ON i.sku=p.product_code
ON DUPLICATE KEY UPDATE presentation_ml=VALUES(presentation_ml),source_reference=VALUES(source_reference);

INSERT INTO items (category_id,base_uom_id,sku,name,item_type,lot_controlled,quality_required,created_by)
SELECT c.id,u.id,m.material_code,m.material_name,'PACKAGING',1,1,@system_admin_id
FROM tmp_core_materials m INNER JOIN item_categories c ON c.code='PACK' INNER JOIN units_of_measure u ON u.code=m.uom_code
ON DUPLICATE KEY UPDATE category_id=VALUES(category_id),base_uom_id=items.base_uom_id,name=VALUES(name),item_type='PACKAGING',lot_controlled=1,quality_required=1,is_active=1;

INSERT INTO items (category_id,base_uom_id,sku,name,item_type,lot_controlled,quality_required,created_by)
SELECT c.id,u.id,r.material_code,r.material_name,r.item_type,1,1,@system_admin_id
FROM tmp_core_raw r INNER JOIN item_categories c ON c.code=IF(r.item_type='RAW_FRUIT','FRUIT','BULK') INNER JOIN units_of_measure u ON u.code=r.uom_code
ON DUPLICATE KEY UPDATE category_id=VALUES(category_id),base_uom_id=items.base_uom_id,name=VALUES(name),item_type=VALUES(item_type),lot_controlled=1,quality_required=1,is_active=1;

-- Corrige la unidad oficial sólo cuando la clave todavía no tiene historial contable.
-- Si ya existen movimientos, conserva su unidad para no reinterpretar cantidades anteriores.
UPDATE items i INNER JOIN tmp_core_pt source ON source.product_code=i.sku INNER JOIN units_of_measure u ON u.code='EA'
SET i.base_uom_id=u.id WHERE NOT EXISTS (SELECT 1 FROM inventory_movement_lines movement WHERE movement.item_id=i.id);
UPDATE items i INNER JOIN tmp_core_materials source ON source.material_code=i.sku INNER JOIN units_of_measure u ON u.code=source.uom_code
SET i.base_uom_id=u.id WHERE NOT EXISTS (SELECT 1 FROM inventory_movement_lines movement WHERE movement.item_id=i.id);
UPDATE items i INNER JOIN tmp_core_raw source ON source.material_code=i.sku INNER JOIN units_of_measure u ON u.code=source.uom_code
SET i.base_uom_id=u.id WHERE NOT EXISTS (SELECT 1 FROM inventory_movement_lines movement WHERE movement.item_id=i.id);

-- Oculta claves antiguas que ya no pertenecen al catálogo oficial sólo cuando no tienen movimientos ni órdenes.
UPDATE items i LEFT JOIN tmp_core_pt p ON p.product_code=i.sku
SET i.is_active=0 WHERE i.item_type='FINISHED_GOOD' AND i.sku LIKE 'PT%' AND p.product_code IS NULL
AND NOT EXISTS (SELECT 1 FROM inventory_movement_lines ml WHERE ml.item_id=i.id)
AND NOT EXISTS (SELECT 1 FROM production_orders po WHERE po.finished_item_id=i.id);

-- BOM de materiales de empaque por producto terminado.
INSERT INTO bills_of_material (bom_code,output_item_id,name,is_active,created_by)
SELECT CONCAT('BOM-ME-',p.product_code),i.id,CONCAT('Materiales de empaque · ',p.product_code),1,@system_admin_id
FROM tmp_core_pt p INNER JOIN items i ON i.sku=p.product_code
ON DUPLICATE KEY UPDATE output_item_id=VALUES(output_item_id),name=VALUES(name),is_active=1;

INSERT INTO bom_versions (bom_id,version_no,output_qty,output_uom_id,effective_from,status,expected_yield_pct,source_reference,approved_by,approved_at)
SELECT b.id,COALESCE((SELECT MAX(existing.version_no)+1 FROM bom_versions existing WHERE existing.bom_id=b.id),1),1,i.base_uom_id,'{year}-01-01','ACTIVE',100,@packaging_reference,@system_admin_id,UTC_TIMESTAMP(6)
FROM tmp_core_pt p INNER JOIN items i ON i.sku=p.product_code INNER JOIN bills_of_material b ON b.bom_code=CONCAT('BOM-ME-',p.product_code)
ON DUPLICATE KEY UPDATE output_qty=1,output_uom_id=VALUES(output_uom_id),effective_from=VALUES(effective_from),effective_to=NULL,status='ACTIVE',approved_by=VALUES(approved_by),approved_at=VALUES(approved_at);

DELETE bc FROM bom_components bc INNER JOIN bom_versions bv ON bv.id=bc.bom_version_id WHERE bv.source_reference=@packaging_reference;
INSERT INTO bom_components (bom_version_id,line_no,component_item_id,component_uom_id,qty_per_output,scrap_pct,is_optional)
SELECT bv.id,c.line_no,i.id,i.base_uom_id,c.qty_per_unit,0,0 FROM tmp_core_pack_components c
INNER JOIN bills_of_material b ON b.bom_code=CONCAT('BOM-ME-',c.product_code)
INNER JOIN bom_versions bv ON bv.bom_id=b.id AND bv.source_reference=@packaging_reference
INNER JOIN items i ON i.sku=c.material_code;

-- Conserva versiones anteriores para trazabilidad, pero evita duplicar necesidades en la planeación nueva.
UPDATE bills_of_material old_bom INNER JOIN items i ON i.id=old_bom.output_item_id INNER JOIN tmp_core_pt p ON p.product_code=i.sku
SET old_bom.is_active=0 WHERE old_bom.bom_code NOT LIKE 'BOM-ME-%' AND old_bom.bom_code NOT LIKE 'BOM-PLAN-%';

-- Formulaciones de aceites versionadas.
INSERT INTO oil_formulations (formulation_code,name,is_active,created_by)
SELECT formulation_code,formulation_name,1,@system_admin_id FROM tmp_core_formulations
ON DUPLICATE KEY UPDATE name=VALUES(name),is_active=1;
INSERT INTO oil_formulation_versions (oil_formulation_id,version_no,effective_from,status,source_reference,approved_by,approved_at)
SELECT f.id,COALESCE((SELECT MAX(existing.version_no)+1 FROM oil_formulation_versions existing WHERE existing.oil_formulation_id=f.id),1),'{year}-01-01','ACTIVE',@formula_reference,@system_admin_id,UTC_TIMESTAMP(6)
FROM tmp_core_formulations source INNER JOIN oil_formulations f ON f.formulation_code=source.formulation_code
ON DUPLICATE KEY UPDATE effective_from=VALUES(effective_from),effective_to=NULL,status='ACTIVE',approved_by=VALUES(approved_by),approved_at=VALUES(approved_at);
DELETE c FROM oil_formulation_components c INNER JOIN oil_formulation_versions v ON v.id=c.oil_formulation_version_id WHERE v.source_reference=@formula_reference;
INSERT INTO oil_formulation_components (oil_formulation_version_id,line_no,raw_item_id,percentage)
SELECT v.id,c.line_no,i.id,c.percentage FROM tmp_core_formula_components c
INNER JOIN oil_formulations f ON f.formulation_code=c.formulation_code
INNER JOIN oil_formulation_versions v ON v.oil_formulation_id=f.id AND v.source_reference=@formula_reference
INNER JOIN items i ON i.sku=c.material_code;

-- Cada PT queda ligado inmediatamente a su BOM de empaque. La formulación de aceite
-- permanece pendiente porque ninguno de los archivos suministrados contiene esa relación.
INSERT INTO finished_product_bom_assignments (finished_item_id,packaging_bom_version_id,effective_from,assignment_status,change_reason,created_by)
SELECT i.id,bv.id,'{year}-01-01','PENDING_FORMULATION','Importación inicial: falta confirmar formulación de aceite',@system_admin_id
FROM tmp_core_pt p INNER JOIN items i ON i.sku=p.product_code
INNER JOIN bills_of_material b ON b.bom_code=CONCAT('BOM-ME-',p.product_code)
INNER JOIN bom_versions bv ON bv.bom_id=b.id AND bv.source_reference=@packaging_reference
WHERE NOT EXISTS (SELECT 1 FROM finished_product_bom_assignments a WHERE a.finished_item_id=i.id);

-- Forecast original y distribución semanal por días laborables.
INSERT INTO forecast_plans (plan_code,name,fiscal_year,status,created_by)
VALUES ('FORECAST-{year}','Planeación de pedidos {year}',{year},'ACTIVE',@system_admin_id)
ON DUPLICATE KEY UPDATE name=VALUES(name),fiscal_year=VALUES(fiscal_year),status='ACTIVE';
SET @forecast_plan_id=(SELECT id FROM forecast_plans WHERE plan_code='FORECAST-{year}' LIMIT 1);
INSERT INTO forecast_versions (forecast_plan_id,version_no,version_type,version_name,source_file_name,source_file_hash,change_reason,is_current,approved_by,approved_at,imported_by)
SELECT @forecast_plan_id,COALESCE(MAX(existing.version_no),0)+1,'ORIGINAL','Catálogo inicial depurado','Forecast_planeacion de pedidos 2026.xlsx',@forecast_hash,'Reingeniería inicial basada en catálogo oficial',0,@system_admin_id,UTC_TIMESTAMP(6),@system_admin_id
FROM forecast_versions existing WHERE existing.forecast_plan_id=@forecast_plan_id
HAVING NOT EXISTS (SELECT 1 FROM forecast_versions same_source WHERE same_source.forecast_plan_id=@forecast_plan_id AND same_source.source_file_hash=@forecast_hash);
SET @forecast_version_id=(SELECT id FROM forecast_versions WHERE forecast_plan_id=@forecast_plan_id AND source_file_hash=@forecast_hash ORDER BY id DESC LIMIT 1);
UPDATE forecast_versions SET is_current=0 WHERE forecast_plan_id=@forecast_plan_id AND id<>@forecast_version_id AND is_current=1;
UPDATE forecast_versions SET is_current=1 WHERE id=@forecast_version_id;
DELETE FROM forecast_lines WHERE forecast_version_id=@forecast_version_id;
INSERT INTO forecast_lines (forecast_version_id,finished_item_id,period_start,forecast_qty,source_row_key)
SELECT @forecast_version_id,i.id,f.period_start,f.forecast_qty,CONCAT(f.product_code,'-FILA-',f.source_row)
FROM tmp_core_forecast_monthly f INNER JOIN items i ON i.sku=f.product_code;
DELETE FROM forecast_weekly_lines WHERE forecast_version_id=@forecast_version_id;
INSERT INTO forecast_weekly_lines (forecast_version_id,finished_item_id,week_start,forecast_qty,allocation_method,source_month,source_row)
SELECT @forecast_version_id,i.id,f.week_start,f.forecast_qty,'WORKDAYS',f.source_month,f.source_row
FROM tmp_core_forecast_weekly f INNER JOIN items i ON i.sku=f.product_code;

DELETE FROM source_import_issues WHERE source_reference=@source_reference;
INSERT INTO source_import_issues (source_reference,severity,issue_code,entity_code,source_row,message)
SELECT source_reference,severity,issue_code,entity_code,source_row,message FROM tmp_core_issues;

INSERT IGNORE INTO schema_migrations (version_no,description)
VALUES ('008',CONCAT('Reingeniería núcleo · fuente ',@source_reference));

COMMIT;

SELECT
  (SELECT COUNT(*) FROM tmp_core_pt) AS productos_terminados,
  (SELECT COUNT(*) FROM tmp_core_materials) AS materiales_empaque,
  (SELECT COUNT(*) FROM tmp_core_pack_components) AS relaciones_empaque,
  (SELECT COUNT(*) FROM tmp_core_formulations) AS formulaciones_aceite,
  (SELECT COUNT(*) FROM tmp_core_forecast_monthly) AS registros_forecast_mensual,
  (SELECT COUNT(*) FROM tmp_core_forecast_weekly) AS registros_forecast_semanal,
  (SELECT COUNT(*) FROM tmp_core_issues WHERE severity='ERROR') AS incidencias_por_corregir;
""",
    ]
    return "\n\n".join(section for section in sections if section)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("product_catalog", type=Path)
    parser.add_argument("forecast", type=Path)
    parser.add_argument("oil_formulas", type=Path)
    parser.add_argument("packaging_bom", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--year", type=int, default=2026)
    args = parser.parse_args()
    paths = [args.product_catalog, args.forecast, args.oil_formulas, args.packaging_bom]
    products, users = load_products(args.product_catalog)
    packaging, materials, packaging_issues = load_packaging(args.packaging_bom, products)
    formulas, raw_materials, formula_issues = load_formulas(args.oil_formulas)
    monthly, weekly, forecast_issues = load_forecast(args.forecast, products, args.year)
    issues = packaging_issues + formula_issues + forecast_issues
    sql = build_sql(paths, products, users, packaging, materials, formulas, raw_materials,
                    monthly, weekly, issues, args.year)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(sql, encoding="utf-8", newline="\n")
    print(f"Generado: {args.output}")
    print(f"Producto terminado oficial: {len(products)}")
    print(f"Materiales de empaque: {len(materials)}")
    print(f"Relaciones de empaque oficiales: {sum(len(rows) for rows in packaging.values())}")
    print(f"Formulaciones de aceite: {len(formulas)}")
    print(f"Forecast mensual importable: {len(monthly)}")
    print(f"Forecast semanal: {len(weekly)}")
    print(f"Incidencias registradas: {len(issues) + 2}")


if __name__ == "__main__":
    main()
